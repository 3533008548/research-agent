"""Persistent, versioned research dossiers outside chat memory.

Markdown remains the source of truth. A dossier uses four stable sections so
the agent can safely patch only what it inspected, while review can cover all
sections without fitting an ever-growing document into one model request.
"""

from __future__ import annotations

import hashlib
import json
import os
import re
import secrets
from datetime import datetime, timezone
from pathlib import Path
from threading import RLock
from typing import Any

from runtime_paths import RuntimePaths


_DOCUMENT_ID = re.compile(r"research-doc-[a-f0-9]{12}\Z")
_LEDGER_ITEM_ID = re.compile(r"ledger-[a-f0-9]{12}\Z")
_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*$")
_store_lock = RLock()
_PLACEHOLDER = "待补充。"

# IDs stay stable even if the visible heading changes in a future template.
RESEARCH_DOSSIER_SECTIONS = (
    ("background", "研究背景与研究现状"),
    ("content_innovation", "研究内容与创新"),
    ("approach_feasibility", "研究方案与可行性"),
    ("outlook_plan", "研究展望与计划"),
)
_SECTION_BY_TITLE = {title: section_id for section_id, title in RESEARCH_DOSSIER_SECTIONS}
_SECTION_TITLES = {section_id: title for section_id, title in RESEARCH_DOSSIER_SECTIONS}
_LEDGER_KINDS = frozenset({"research_question", "hypothesis", "innovation_candidate", "decision"})
_LEDGER_STATUSES = frozenset({"open", "hypothesis", "supported", "contested", "rejected", "decision"})
_EVIDENCE_RELATIONS = frozenset({"supports", "contradicts", "conditions", "inspiration"})


class ResearchDocumentStore:
    """Store user-owned dossiers with atomic section patches and snapshots."""

    def __init__(self, paths: RuntimePaths) -> None:
        self.paths = paths
        self.paths.ensure_initialized()
        self.directory = paths.research_documents_dir

    def save(self, title: str, content: str, document_id: str = "") -> dict[str, Any]:
        """Create a dossier; complete replacement of an existing one is unsafe."""
        if str(document_id or "").strip():
            raise ValueError("已有科研档案不能整篇覆盖；请读取目标章节后提交局部补丁")
        title = _clean_title(title)
        content = str(content or "").strip()
        if not title:
            raise ValueError("文档标题不能为空")
        if not content:
            raise ValueError("文档内容不能为空")

        with _store_lock:
            document_id = f"research-doc-{secrets.token_hex(6)}"
            now = _now()
            normalized = _normalise_dossier(title, content)
            metadata = {
                "version": 2,
                "template": "research-dossier/v1",
                "document_id": document_id,
                "title": title[:160],
                "summary": _summary(normalized),
                "created_at": now,
                "updated_at": now,
                "revision": 1,
                "markdown_file": "document.md",
                "docx_file": "document.docx",
            }
            directory = self.directory / document_id
            directory.mkdir(parents=True, exist_ok=False)
            self._write_current(directory, metadata, normalized)
            self._write_ledger(directory, _empty_ledger(metadata))
            return self._with_paths(metadata)

    def list(self, limit: int = 50) -> list[dict[str, Any]]:
        """List document metadata without reading whole Markdown bodies."""
        limit = max(1, min(int(limit), 200))
        with _store_lock:
            documents = [
                self._with_paths(metadata)
                for metadata_path in self.directory.glob("research-doc-*/metadata.json")
                if (metadata := self._read_metadata_file(metadata_path)) is not None
            ]
        return sorted(documents, key=lambda item: str(item["updated_at"]), reverse=True)[:limit]

    def search(self, query: str, limit: int = 5) -> list[dict[str, Any]]:
        """Search titles and Markdown bodies; empty query lists recent docs."""
        query = " ".join(str(query or "").split()).strip()
        query_folded = query.casefold()
        limit = max(1, min(int(limit), 10))
        matches: list[tuple[int, dict[str, Any]]] = []
        for metadata in self.list(limit=200):
            content = self._read_content(metadata)
            title = str(metadata["title"])
            if not query:
                matches.append((0, {**metadata, "excerpt": _summary(content)}))
                continue
            title_hits = title.casefold().count(query_folded)
            content_hits = content.casefold().count(query_folded)
            if title_hits or content_hits:
                matches.append((title_hits * 8 + content_hits, {
                    **metadata,
                    "excerpt": _excerpt(content, query_folded),
                }))
        matches.sort(key=lambda item: (item[0], item[1]["updated_at"]), reverse=True)
        return [item for _, item in matches[:limit]]

    def read(self, document_id: str) -> dict[str, Any] | None:
        """Read one dossier for the UI, never as authorization for blind writes."""
        with _store_lock:
            metadata = self._load_metadata(document_id)
            if metadata is None:
                return None
            content = self._read_content(metadata)
            return {
                **self._with_paths(metadata),
                "content": content,
                "sections": _section_descriptors(content),
                "ledger": self._ledger_payload(metadata, content),
            }

    def list_sections(self, document_id: str) -> dict[str, Any] | None:
        """Return compact, hashed section metadata for planning a patch."""
        with _store_lock:
            metadata = self._load_metadata(document_id)
            if metadata is None:
                return None
            content = self._read_content(metadata)
            return {**self._with_paths(metadata), "sections": _section_descriptors(content)}

    def read_section(self, document_id: str, section_id: str) -> dict[str, Any] | None:
        """Read exactly one named section along with its optimistic hash."""
        with _store_lock:
            metadata = self._load_metadata(document_id)
            if metadata is None:
                return None
            section_id = str(section_id or "").strip()
            sections = _split_sections(self._read_content(metadata))
            if section_id not in sections:
                raise ValueError("未知科研档案章节")
            body = sections[section_id]
            return {
                **self._with_paths(metadata),
                "section_id": section_id,
                "heading": _SECTION_TITLES[section_id],
                "content": body,
                "content_hash": _content_hash(body),
            }

    def read_ledger(self, document_id: str) -> dict[str, Any] | None:
        """Read the dossier's evidence and hypothesis ledger with link freshness."""
        with _store_lock:
            metadata = self._load_metadata(document_id)
            if metadata is None:
                return None
            content = self._read_content(metadata)
            return self._ledger_payload(metadata, content)

    def apply_ledger_patch(
        self,
        document_id: str,
        *,
        base_revision: int,
        operations: list[dict[str, Any]],
    ) -> dict[str, Any]:
        """Update explicit research claims without touching dossier prose.

        Each upsert is tied to the hash of its target section.  This prevents a
        stale interpretation of an earlier section from being represented as a
        current research claim after the prose has changed.
        """
        with _store_lock:
            metadata = self._load_metadata(document_id)
            if metadata is None:
                raise ValueError("未找到要更新账本的研究档案")
            content = self._read_content(metadata)
            sections = _split_sections(content)
            ledger = self._read_ledger(metadata)
            try:
                requested_revision = int(base_revision)
            except (TypeError, ValueError) as exc:
                raise ValueError("账本更新必须携带当前账本修订号") from exc
            if requested_revision != int(ledger.get("ledger_revision") or 0):
                raise ValueError("证据与假设账本已更新，请重新读取后再修改")
            if not isinstance(operations, list) or not operations:
                raise ValueError("账本补丁至少包含一个操作")
            if len(operations) > 20:
                raise ValueError("一次最多修改 20 条账本项")

            items = {str(item["item_id"]): dict(item) for item in ledger.get("items") or []}
            changes: list[dict[str, str]] = []
            for raw in operations:
                if not isinstance(raw, dict):
                    raise ValueError("账本操作格式错误")
                operation = str(raw.get("op") or "").strip()
                item_id = str(raw.get("item_id") or "").strip()
                if operation == "delete_item":
                    if item_id not in items:
                        raise ValueError("要删除的账本项不存在")
                    removed = items.pop(item_id)
                    changes.append({"action": "删除", "item_id": item_id, "statement": str(removed.get("statement") or "")})
                    continue
                if operation != "upsert_item":
                    raise ValueError("账本仅支持 upsert_item 或 delete_item")
                item = _clean_ledger_item(raw, sections, item_id=item_id)
                items[item["item_id"]] = item
                changes.append({"action": "更新" if item_id else "新增", "item_id": item["item_id"], "statement": item["statement"]})

            updated = {
                **ledger,
                "version": 1,
                "document_id": str(metadata["document_id"]),
                "title": str(metadata["title"]),
                "ledger_revision": int(ledger.get("ledger_revision") or 0) + 1,
                "updated_at": _now(),
                "items": sorted(items.values(), key=lambda item: (str(item.get("section_id") or ""), str(item.get("item_id") or ""))),
            }
            self._write_ledger(self.directory / str(metadata["document_id"]), updated)
            payload = self._ledger_payload(metadata, content, ledger=updated)
            return {**payload, "changes": changes}

    def apply_patch(
        self,
        document_id: str,
        *,
        base_revision: int,
        operations: list[dict[str, Any]],
    ) -> dict[str, Any]:
        """Apply checked section replacements/appends; never replace an unseen file."""
        with _store_lock:
            metadata = self._load_metadata(document_id)
            if metadata is None:
                raise ValueError("未找到要更新的研究档案")
            try:
                requested_revision = int(base_revision)
            except (TypeError, ValueError) as exc:
                raise ValueError("补丁必须携带当前修订号") from exc
            if requested_revision != int(metadata.get("revision") or 0):
                raise ValueError("科研档案已更新，请重新读取目标章节后再修改")
            if not isinstance(operations, list) or not operations:
                raise ValueError("补丁至少包含一个章节操作")
            if len(operations) > len(RESEARCH_DOSSIER_SECTIONS):
                raise ValueError("一次最多修改四个固定章节")

            current = self._read_content(metadata)
            sections = _split_sections(current)
            changed: list[dict[str, str]] = []
            touched: set[str] = set()
            for raw in operations:
                if not isinstance(raw, dict):
                    raise ValueError("补丁操作格式错误")
                operation = str(raw.get("op") or "").strip()
                section_id = str(raw.get("section_id") or "").strip()
                expected_hash = str(raw.get("expected_hash") or "").strip()
                if operation not in {"replace_section", "append_to_section"}:
                    raise ValueError("仅支持 replace_section 或 append_to_section")
                if section_id not in sections:
                    raise ValueError("补丁指定了未知科研档案章节")
                if section_id in touched:
                    raise ValueError("同一章节一次只能修改一次")
                if not expected_hash or expected_hash != _content_hash(sections[section_id]):
                    raise ValueError(f"章节“{_SECTION_TITLES[section_id]}”已变化，请重新读取后再修改")
                replacement = _clean_section_body(raw.get("content"))
                if operation == "replace_section":
                    sections[section_id] = replacement
                    action = "替换"
                else:
                    prefix = "" if sections[section_id] == _PLACEHOLDER else "\n\n"
                    sections[section_id] = sections[section_id] + prefix + replacement
                    action = "追加"
                changed.append({"section_id": section_id, "heading": _SECTION_TITLES[section_id], "action": action})
                touched.add(section_id)

            updated_content = _render_dossier(str(metadata["title"]), sections)
            updated = self._next_metadata(metadata, updated_content)
            directory = self.directory / str(metadata["document_id"])
            self._snapshot_current(directory, metadata, current)
            self._write_current(directory, updated, updated_content)
            return {**self._with_paths(updated), "changes": changed, "sections": _section_descriptors(updated_content)}

    def list_versions(self, document_id: str, limit: int = 20) -> list[dict[str, Any]]:
        """List restorable snapshots plus the current revision without bodies."""
        limit = max(1, min(int(limit), 50))
        with _store_lock:
            metadata = self._load_metadata(document_id)
            if metadata is None:
                return []
            versions = [{"revision": int(metadata["revision"]), "updated_at": str(metadata["updated_at"]), "current": True}]
            history = self.directory / str(metadata["document_id"]) / "history"
            for snapshot in history.glob("revision-*.json"):
                previous = self._read_metadata_file(snapshot, require_current_file=False)
                if previous is not None:
                    versions.append({"revision": int(previous.get("revision") or 0), "updated_at": str(previous.get("updated_at") or ""), "current": False})
            versions.sort(key=lambda item: int(item["revision"]), reverse=True)
            return versions[:limit]

    def restore_version(self, document_id: str, *, base_revision: int, revision: int) -> dict[str, Any]:
        """Restore a historical body as a new revision instead of rewinding history."""
        with _store_lock:
            metadata = self._load_metadata(document_id)
            if metadata is None:
                raise ValueError("未找到要恢复的研究档案")
            if int(base_revision) != int(metadata.get("revision") or 0):
                raise ValueError("科研档案已更新，请刷新后再恢复历史版本")
            try:
                target = int(revision)
            except (TypeError, ValueError) as exc:
                raise ValueError("历史版本号无效") from exc
            snapshot = self.directory / str(metadata["document_id"]) / "history" / f"revision-{target:04d}.md"
            try:
                restored = snapshot.read_text(encoding="utf-8").strip()
            except OSError as exc:
                raise ValueError("未找到可恢复的历史版本") from exc
            current = self._read_content(metadata)
            updated = self._next_metadata(metadata, restored)
            directory = self.directory / str(metadata["document_id"])
            self._snapshot_current(directory, metadata, current)
            self._write_current(directory, updated, restored)
            return {**self._with_paths(updated), "restored_from_revision": target}

    def _load_metadata(self, document_id: str) -> dict[str, Any] | None:
        document_id = str(document_id or "").strip()
        if not _DOCUMENT_ID.fullmatch(document_id):
            return None
        return self._read_metadata_file(self.directory / document_id / "metadata.json")

    @staticmethod
    def _read_metadata_file(path: Path, *, require_current_file: bool = True) -> dict[str, Any] | None:
        try:
            metadata = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            return None
        if not _DOCUMENT_ID.fullmatch(str(metadata.get("document_id") or "")):
            return None
        if not str(metadata.get("title") or "").strip():
            return None
        if require_current_file and str(metadata.get("markdown_file") or "document.md") != "document.md":
            return None
        return metadata

    def _read_content(self, metadata: dict[str, Any]) -> str:
        path = self.paths.safe_child(
            self.directory / str(metadata["document_id"]),
            str(metadata.get("markdown_file") or "document.md"),
        )
        try:
            return path.read_text(encoding="utf-8").strip()
        except OSError:
            return ""

    def _read_ledger(self, metadata: dict[str, Any]) -> dict[str, Any]:
        path = self.directory / str(metadata["document_id"]) / "research_ledger.json"
        try:
            raw = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            return _empty_ledger(metadata)
        if not isinstance(raw, dict) or str(raw.get("document_id") or "") != str(metadata["document_id"]):
            return _empty_ledger(metadata)
        items = raw.get("items") if isinstance(raw.get("items"), list) else []
        return {
            "version": 1,
            "document_id": str(metadata["document_id"]),
            "title": str(metadata["title"]),
            "ledger_revision": _non_negative_int(raw.get("ledger_revision")),
            "created_at": str(raw.get("created_at") or metadata.get("created_at") or _now()),
            "updated_at": str(raw.get("updated_at") or metadata.get("updated_at") or _now()),
            "items": [dict(item) for item in items if isinstance(item, dict) and _LEDGER_ITEM_ID.fullmatch(str(item.get("item_id") or ""))],
        }

    def _ledger_payload(
        self,
        metadata: dict[str, Any],
        content: str,
        *,
        ledger: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        current = ledger or self._read_ledger(metadata)
        sections = _split_sections(content)
        items = [_hydrate_ledger_item(item, sections) for item in current.get("items") or []]
        return {
            "document_id": str(metadata["document_id"]),
            "title": str(metadata["title"]),
            "document_revision": int(metadata.get("revision") or 0),
            "ledger_revision": int(current.get("ledger_revision") or 0),
            "updated_at": str(current.get("updated_at") or metadata.get("updated_at") or ""),
            "items": items,
            "summary": _ledger_summary(items),
        }

    def _with_paths(self, metadata: dict[str, Any]) -> dict[str, Any]:
        directory = self.directory / str(metadata["document_id"])
        return {
            **metadata,
            "markdown_path": str(directory / str(metadata.get("markdown_file") or "document.md")),
            "docx_path": str(directory / str(metadata.get("docx_file") or "document.docx")),
            "ledger_path": str(directory / "research_ledger.json"),
        }

    @staticmethod
    def _next_metadata(metadata: dict[str, Any], content: str) -> dict[str, Any]:
        return {
            **metadata,
            "version": 2,
            "template": "research-dossier/v1",
            "summary": _summary(content),
            "updated_at": _now(),
            "revision": int(metadata.get("revision") or 0) + 1,
        }

    @staticmethod
    def _snapshot_current(directory: Path, metadata: dict[str, Any], content: str) -> None:
        history = directory / "history"
        history.mkdir(parents=True, exist_ok=True)
        revision = int(metadata.get("revision") or 0)
        _write_text(history / f"revision-{revision:04d}.md", content.rstrip() + "\n")
        _write_json(history / f"revision-{revision:04d}.json", metadata)

    def _write_current(self, directory: Path, metadata: dict[str, Any], content: str) -> None:
        markdown_path = directory / str(metadata.get("markdown_file") or "document.md")
        docx_path = directory / str(metadata.get("docx_file") or "document.docx")
        staged_markdown = directory / "document.next.md"
        staged_docx = directory / "document.next.docx"
        _write_text(staged_markdown, content.rstrip() + "\n")
        self._write_docx(staged_docx, str(metadata["title"]), content)
        os.replace(staged_markdown, markdown_path)
        os.replace(staged_docx, docx_path)
        _write_json(directory / "metadata.json", metadata)

    @staticmethod
    def _write_ledger(directory: Path, ledger: dict[str, Any]) -> None:
        _write_json(directory / "research_ledger.json", ledger)

    @staticmethod
    def _write_docx(path: Path, title: str, content: str) -> None:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt

        document = Document()
        page = document.sections[0]
        page.top_margin = Inches(0.8)
        page.bottom_margin = Inches(0.8)
        page.left_margin = Inches(0.85)
        page.right_margin = Inches(0.85)
        normal = document.styles["Normal"]
        normal.font.name = "Aptos"
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        normal.font.size = Pt(11)

        title_paragraph = document.add_paragraph()
        title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title_paragraph.add_run(title)
        title_run.bold = True
        title_run.font.name = "Aptos Display"
        title_run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        title_run.font.size = Pt(20)

        first_heading = True
        for raw_line in content.splitlines():
            line = raw_line.strip()
            if not line:
                continue
            heading = re.match(r"^(#{1,3})\s+(.+)$", line)
            if heading:
                heading_text = heading.group(2).strip()
                if first_heading and heading_text == title:
                    first_heading = False
                    continue
                first_heading = False
                document.add_heading(heading_text, level=len(heading.group(1)))
            else:
                first_heading = False
                bullet = re.match(r"^[-*]\s+(.+)$", line)
                numbered = re.match(r"^\d+[.)]\s+(.+)$", line)
                if bullet:
                    document.add_paragraph(bullet.group(1), style="List Bullet")
                elif numbered:
                    document.add_paragraph(numbered.group(1), style="List Number")
                else:
                    document.add_paragraph(line)

        temporary = path.with_suffix(path.suffix + ".tmp")
        document.core_properties.title = title
        document.core_properties.author = "Research Agent"
        document.save(str(temporary))
        os.replace(temporary, path)


def _clean_title(value: object) -> str:
    return " ".join(str(value or "").split()).strip()


def _normalise_dossier(title: str, content: str) -> str:
    """Apply the fixed layout without discarding unstructured draft text."""
    parsed: dict[str, list[str]] = {section_id: [] for section_id, _ in RESEARCH_DOSSIER_SECTIONS}
    preamble: list[str] = []
    active: str | None = None
    for raw_line in str(content or "").replace("\r\n", "\n").split("\n"):
        heading = _HEADING.match(raw_line.strip())
        heading_title = heading.group(2).strip() if heading else ""
        section_id = _SECTION_BY_TITLE.get(heading_title)
        if section_id and len(heading.group(1)) <= 2:
            active = section_id
            continue
        if heading and heading_title == title and len(heading.group(1)) == 1:
            continue
        (parsed[active] if active else preamble).append(raw_line)
    if preamble:
        parsed["background"] = preamble + parsed["background"]
    return _render_dossier(title, {
        section_id: _clean_section_body("\n".join(lines), allow_placeholder=True)
        for section_id, lines in parsed.items()
    })


def _split_sections(content: str) -> dict[str, str]:
    """Read canonical bodies, normalising legacy documents on the fly."""
    title = _extract_title(content) or "科研档案"
    normalized = _normalise_dossier(title, content)
    sections: dict[str, list[str]] = {section_id: [] for section_id, _ in RESEARCH_DOSSIER_SECTIONS}
    active: str | None = None
    for raw_line in normalized.splitlines():
        heading = _HEADING.match(raw_line.strip())
        if heading and len(heading.group(1)) == 2:
            section_id = _SECTION_BY_TITLE.get(heading.group(2).strip())
            if section_id:
                active = section_id
                continue
        if heading and len(heading.group(1)) == 1:
            continue
        if active:
            sections[active].append(raw_line)
    return {
        section_id: _clean_section_body("\n".join(lines), allow_placeholder=True)
        for section_id, lines in sections.items()
    }


def _render_dossier(title: str, sections: dict[str, str]) -> str:
    blocks = [f"# {title}"]
    for section_id, heading in RESEARCH_DOSSIER_SECTIONS:
        blocks.append(f"## {heading}\n\n{_clean_section_body(sections.get(section_id, ''), allow_placeholder=True)}")
    return "\n\n".join(blocks).strip()


def _extract_title(content: str) -> str:
    for raw_line in str(content or "").splitlines():
        heading = _HEADING.match(raw_line.strip())
        if heading and len(heading.group(1)) == 1:
            return _clean_title(heading.group(2))
    return ""


def _clean_section_body(value: object, *, allow_placeholder: bool = False) -> str:
    body = str(value or "").strip()
    if not body:
        return _PLACEHOLDER if allow_placeholder else _PLACEHOLDER
    for _section_id, heading in RESEARCH_DOSSIER_SECTIONS:
        if re.search(rf"^##\s+{re.escape(heading)}\s*$", body, flags=re.MULTILINE):
            raise ValueError("章节正文不能包含固定章节标题")
    return body


def _section_descriptors(content: str) -> list[dict[str, Any]]:
    sections = _split_sections(content)
    return [{
        "section_id": section_id,
        "heading": heading,
        "content_length": len(sections[section_id]),
        "content_hash": _content_hash(sections[section_id]),
        "summary": _summary(sections[section_id], limit=240),
    } for section_id, heading in RESEARCH_DOSSIER_SECTIONS]


def _content_hash(content: str) -> str:
    return hashlib.sha256(str(content).encode("utf-8")).hexdigest()[:16]


def _empty_ledger(metadata: dict[str, Any]) -> dict[str, Any]:
    now = str(metadata.get("created_at") or _now())
    return {
        "version": 1,
        "document_id": str(metadata["document_id"]),
        "title": str(metadata["title"]),
        "ledger_revision": 0,
        "created_at": now,
        "updated_at": str(metadata.get("updated_at") or now),
        "items": [],
    }


def _clean_ledger_item(
    raw: dict[str, Any],
    sections: dict[str, str],
    *,
    item_id: str,
) -> dict[str, Any]:
    resolved_id = item_id or f"ledger-{secrets.token_hex(6)}"
    if not _LEDGER_ITEM_ID.fullmatch(resolved_id):
        raise ValueError("账本项 ID 格式无效")
    section_id = str(raw.get("section_id") or "").strip()
    if section_id not in sections:
        raise ValueError("账本项必须关联一个固定科研档案章节")
    expected_section_hash = str(raw.get("expected_section_hash") or "").strip()
    if expected_section_hash != _content_hash(sections[section_id]):
        raise ValueError(f"章节“{_SECTION_TITLES[section_id]}”已变化，请重新读取后再记录账本")
    kind = str(raw.get("kind") or "").strip()
    status = str(raw.get("status") or "").strip()
    if kind not in _LEDGER_KINDS:
        raise ValueError("账本项类型无效")
    if status not in _LEDGER_STATUSES:
        raise ValueError("账本项认识状态无效")
    statement = _bounded_text(raw.get("statement"), "账本项陈述", 800)
    falsification = _optional_bounded_text(raw.get("falsification"), "可证伪条件", 500)
    evidence_raw = raw.get("evidence") or []
    if not isinstance(evidence_raw, list) or len(evidence_raw) > 12:
        raise ValueError("每条账本项最多关联 12 条证据")
    evidence = [_clean_ledger_evidence(item) for item in evidence_raw]
    if status == "supported" and not evidence:
        raise ValueError("标记为 supported 的账本项至少需要一条可定位证据")
    return {
        "item_id": resolved_id,
        "section_id": section_id,
        "section_hash": expected_section_hash,
        "kind": kind,
        "status": status,
        "statement": statement,
        "falsification": falsification,
        "evidence": evidence,
        "updated_at": _now(),
    }


def _clean_ledger_evidence(raw: object) -> dict[str, Any]:
    if not isinstance(raw, dict):
        raise ValueError("账本证据格式错误")
    paper_id = _optional_bounded_text(raw.get("paper_id"), "论文 ID", 160)
    title = _optional_bounded_text(raw.get("title"), "论文标题", 300)
    if not paper_id and not title:
        raise ValueError("账本证据至少需要论文 ID 或标题")
    relation = str(raw.get("relation") or "").strip()
    if relation not in _EVIDENCE_RELATIONS:
        raise ValueError("账本证据关系无效")
    page = raw.get("page")
    if page not in (None, ""):
        try:
            page = int(page)
        except (TypeError, ValueError) as exc:
            raise ValueError("账本证据页码无效") from exc
        if page < 1:
            raise ValueError("账本证据页码必须大于 0")
    else:
        page = None
    chunk_index = raw.get("chunk_index")
    if chunk_index not in (None, ""):
        try:
            chunk_index = int(chunk_index)
        except (TypeError, ValueError) as exc:
            raise ValueError("账本证据切块编号无效") from exc
        if chunk_index < 0:
            raise ValueError("账本证据切块编号不能为负数")
    else:
        chunk_index = None
    return {
        "paper_id": paper_id,
        "title": title,
        "page": page,
        "chunk_index": chunk_index,
        "relation": relation,
        "note": _optional_bounded_text(raw.get("note"), "证据说明", 500),
    }


def _hydrate_ledger_item(item: dict[str, Any], sections: dict[str, str]) -> dict[str, Any]:
    section_id = str(item.get("section_id") or "")
    current_hash = _content_hash(sections[section_id]) if section_id in sections else ""
    evidence = item.get("evidence") if isinstance(item.get("evidence"), list) else []
    return {
        "item_id": str(item.get("item_id") or ""),
        "section_id": section_id,
        "heading": _SECTION_TITLES.get(section_id, "未知章节"),
        "section_hash": str(item.get("section_hash") or ""),
        "section_current": bool(current_hash and str(item.get("section_hash") or "") == current_hash),
        "kind": str(item.get("kind") or ""),
        "status": str(item.get("status") or ""),
        "statement": str(item.get("statement") or ""),
        "falsification": str(item.get("falsification") or ""),
        "evidence": [dict(source) for source in evidence if isinstance(source, dict)],
        "updated_at": str(item.get("updated_at") or ""),
    }


def _ledger_summary(items: list[dict[str, Any]]) -> dict[str, int]:
    return {
        "items": len(items),
        "supported": sum(1 for item in items if item.get("status") == "supported"),
        "contested": sum(1 for item in items if item.get("status") == "contested"),
        "stale_links": sum(1 for item in items if not item.get("section_current", False)),
    }


def _bounded_text(value: object, label: str, limit: int) -> str:
    text = " ".join(str(value or "").split()).strip()
    if not text:
        raise ValueError(f"{label}不能为空")
    if len(text) > limit:
        raise ValueError(f"{label}不能超过 {limit} 个字符")
    return text


def _optional_bounded_text(value: object, label: str, limit: int) -> str:
    text = " ".join(str(value or "").split()).strip()
    if len(text) > limit:
        raise ValueError(f"{label}不能超过 {limit} 个字符")
    return text


def _non_negative_int(value: object) -> int:
    try:
        return max(0, int(value))
    except (TypeError, ValueError):
        return 0


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _summary(content: str, limit: int = 240) -> str:
    compact = " ".join(line.strip("#> -*\t ") for line in content.splitlines() if line.strip())
    return compact[:limit].strip()


def _excerpt(content: str, query: str, radius: int = 180) -> str:
    compact = " ".join(content.split())
    index = compact.casefold().find(query)
    if index < 0:
        return _summary(compact)
    start = max(0, index - radius)
    end = min(len(compact), index + len(query) + radius)
    return ("…" if start else "") + compact[start:end].strip() + ("…" if end < len(compact) else "")


def _write_text(path: Path, text: str) -> None:
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(text, encoding="utf-8")
    os.replace(temporary, path)


def _write_json(path: Path, value: dict[str, Any]) -> None:
    _write_text(path, json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True) + "\n")
